from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "least-privilege-connector-accounts.md"
DOCX_OUT = ROOT / "least-privilege-connector-accounts.docx"
XLSX_OUT = ROOT / "least-privilege-connector-accounts.xlsx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(8.5)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_doc(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(22)
    styles["Title"].font.bold = True
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 121)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 2"].font.bold = True
    styles["Heading 2"].font.color.rgb = RGBColor(31, 78, 121)
    styles["Heading 3"].font.name = "Arial"
    styles["Heading 3"].font.size = Pt(11)
    styles["Heading 3"].font.bold = True
    styles["Heading 3"].font.color.rgb = RGBColor(68, 68, 68)


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            set_cell_text(cell, value.strip(), bold=(r_idx == 0), color="FFFFFF" if r_idx == 0 else None)
            if r_idx == 0:
                set_cell_shading(cell, "1F4E79")
            elif r_idx % 2 == 0:
                set_cell_shading(cell, "F3F6FA")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
    document.add_paragraph()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip()
        cells = [c.strip() for c in raw.strip("|").split("|")]
        if not all(re.fullmatch(r"[-: ]+", c) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def build_docx() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    document = Document()
    style_doc(document)

    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Least-Privilege Discovery Accounts by Connector Type")

    sub = document.add_paragraph()
    sub.add_run("Account Discovery Tool production guide").bold = True
    sub.add_run(f" | Generated {datetime.now().strftime('%Y-%m-%d')}")

    in_code = False
    code_lang = ""
    code_lines: list[str] = []
    i = 1
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lang = stripped.strip("`").strip() or "text"
                code_lines = []
            else:
                p = document.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.18)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Courier New"
                run.font.size = Pt(7.5)
                run.font.color.rgb = RGBColor(40, 40, 40)
                in_code = False
                code_lang = ""
                code_lines = []
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if stripped.startswith("|"):
            table_rows, i = parse_table(lines, i)
            add_markdown_table(document, table_rows)
            continue

        if stripped.startswith("# "):
            i += 1
            continue
        if stripped.startswith("## "):
            document.add_heading(stripped[3:], level=1)
        elif stripped.startswith("### "):
            document.add_heading(stripped[4:], level=2)
        elif stripped.startswith("- "):
            p = document.add_paragraph(style="List Bullet")
            p.add_run(stripped[2:])
        elif re.match(r"^\d+\. ", stripped):
            p = document.add_paragraph(style="List Number")
            p.add_run(re.sub(r"^\d+\. ", "", stripped))
        else:
            p = document.add_paragraph()
            p.add_run(stripped)
        i += 1

    document.save(DOCX_OUT)


def parse_summary_matrix(md: str) -> list[list[str]]:
    marker = "## Summary Matrix"
    start = md.index(marker)
    chunk = md[start:].split("\n## ", 1)[0]
    lines = chunk.splitlines()
    for idx, line in enumerate(lines):
        if line.strip().startswith("|"):
            rows, _ = parse_table(lines, idx)
            return rows
    return []


def parse_commands(md: str) -> list[list[str]]:
    appendix = md.split("## Appendix A: Account Creation Command Examples", 1)[1]
    rows: list[list[str]] = []
    current_platform = ""
    current_text: list[str] = []
    in_code = False
    lang = ""
    code: list[str] = []

    for line in appendix.splitlines():
        if line.startswith("### "):
            current_platform = line[4:].strip()
            current_text = []
            continue
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                lang = line.strip().strip("`") or "text"
                code = []
            else:
                description = " ".join(x.strip() for x in current_text if x.strip())[-240:]
                rows.append([current_platform, description, lang, "\n".join(code)])
                in_code = False
                lang = ""
                code = []
            continue
        if in_code:
            code.append(line)
        elif current_platform and line.strip():
            current_text.append(line)
    return rows


def parse_checklist(md: str) -> list[list[str]]:
    chunk = md.split("## Production Approval Checklist", 1)[1].split("## Appendix A", 1)[0]
    rows = [["No.", "Production Approval Check"]]
    n = 1
    for line in chunk.splitlines():
        if line.strip().startswith("- "):
            rows.append([n, line.strip()[2:]])
            n += 1
    return rows


def add_sheet_table(ws, rows: list[list], table_name: str) -> None:
    for row in rows:
        ws.append(row)
    if not rows:
        return
    end_col = get_column_letter(len(rows[0]))
    end_row = len(rows)
    tab = Table(displayName=table_name, ref=f"A1:{end_col}{end_row}")
    style = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    tab.tableStyleInfo = style
    ws.add_table(tab)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{end_col}{end_row}"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    thin = Side(style="thin", color="D9E2EC")
    for row in ws.iter_rows(min_row=1, max_row=end_row, min_col=1, max_col=len(rows[0])):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if cell.row == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            else:
                cell.font = Font(size=10)

    for col_idx in range(1, len(rows[0]) + 1):
        values = [str(ws.cell(row=r, column=col_idx).value or "") for r in range(1, min(end_row, 30) + 1)]
        width = min(max(max((len(v.splitlines()[0]) if v else 0) for v in values) + 3, 12), 70)
        if rows[0][col_idx - 1] in ("Command", "Required / Avoided Privileges"):
            width = 78
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def build_xlsx() -> None:
    md = SOURCE.read_text(encoding="utf-8")
    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    overview = [
        ["Field", "Value"],
        ["Guide", "Least-Privilege Discovery Accounts by Connector Type"],
        ["Application", "Account Discovery Tool"],
        ["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")],
        ["Purpose", "Connector-specific production account requirements and setup commands"],
        ["Source", str(SOURCE.name)],
    ]
    add_sheet_table(ws, overview, "OverviewTable")

    matrix_rows = parse_summary_matrix(md)
    ws = wb.create_sheet("Connector Matrix")
    add_sheet_table(ws, matrix_rows, "ConnectorMatrix")

    command_rows = [["Platform / Database", "Command Purpose", "Command Type", "Command"]]
    command_rows.extend(parse_commands(md))
    ws = wb.create_sheet("Creation Commands")
    add_sheet_table(ws, command_rows, "CreationCommands")
    ws.column_dimensions["D"].width = 95

    checklist_rows = parse_checklist(md)
    ws = wb.create_sheet("Approval Checklist")
    add_sheet_table(ws, checklist_rows, "ApprovalChecklist")
    ws.column_dimensions["B"].width = 100

    avoid_rows = [
        ["Connector", "Avoid Granting"],
        ["Windows", "Domain Admins, Enterprise Admins, local Administrators by default. Use JEA or documented host-scoped exception when complete evidence requires elevation."],
        ["Linux/Unix SSH", "Broad sudo, unrestricted root shell, write access to account files, general command execution beyond read-only inventory commands."],
        ["MySQL / MariaDB", "SUPER, CREATE USER, GRANT OPTION, ALL PRIVILEGES, write privileges on application schemas."],
        ["MSSQL", "sysadmin, securityadmin, serveradmin, db_owner, ALTER ANY LOGIN, CONTROL SERVER."],
        ["MongoDB", "root, userAdminAnyDatabase, dbAdminAnyDatabase, clusterAdmin, readWriteAnyDatabase, backup unless approved."],
        ["Oracle DB", "DBA, SYSDBA, SYSOPER, SYSBACKUP, SYSKM, GRANT ANY PRIVILEGE, CREATE USER."],
        ["PostgreSQL", "SUPERUSER, CREATEROLE, CREATEDB, REPLICATION, pg_read_server_files, pg_write_server_files, pg_execute_server_program."],
        ["Redis", "+@all, +@write, broad +@admin, FLUSHALL, CONFIG SET, SHUTDOWN, MODULE, EVAL."],
    ]
    ws = wb.create_sheet("Do Not Grant")
    add_sheet_table(ws, avoid_rows, "DoNotGrant")
    ws.column_dimensions["B"].width = 110

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows():
            ws.row_dimensions[row[0].row].height = 24 if row[0].row == 1 else 48

    wb.save(XLSX_OUT)
    load_workbook(XLSX_OUT).close()


if __name__ == "__main__":
    build_docx()
    build_xlsx()
    print(DOCX_OUT)
    print(XLSX_OUT)
