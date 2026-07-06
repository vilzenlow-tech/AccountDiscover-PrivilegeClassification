"""Export endpoints (CSV, Excel) with audit trail."""
from __future__ import annotations

import io
from datetime import UTC, datetime

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.db import get_db
from app.models.account import Account
from app.models.asset import Asset
from app.models.enums import ActivityStatus, EnabledStatus, PrivilegeClass
from app.services.account_context import domain_for_account, origin_for_account, schema_name_for_account
from app.security import Principal, require_roles
from app.services.activity import housekeeping_risk
from app.services.audit import log_action

router = APIRouter(prefix="/exports", tags=["exports"])

# Safety cap: exports are paginated to avoid full-database exfiltration in one request.
_EXPORT_MAX_ROWS = 50_000

PRIVILEGED_CLASSES = [
    PrivilegeClass.full_admin,
    PrivilegeClass.admin_equivalent,
    PrivilegeClass.operator_high_impact,
    PrivilegeClass.delegated_admin,
    PrivilegeClass.privileged_service,
    PrivilegeClass.dormant_privileged,
]


def _accounts_to_rows(accounts: list[Account]) -> list[dict]:
    return [
        {
            "account_id": str(a.id),
            "account_name": a.account_name,
            "hostname": a.asset.hostname if a.asset else "",
            "ip_address": (a.asset.ip_address or "") if a.asset else "",
            "activity_status": a.activity_status.value,
            "password_last_changed": a.password_last_changed.isoformat() if a.password_last_changed else "",
            "never_logged_in": a.never_logged_in if a.never_logged_in is not None else "",
            "collection_mode": a.collection_mode or "",
            "asset_id": str(a.asset_id),
            "platform": a.platform.value,
            "schema_name": schema_name_for_account(a) or "",
            "source_type": a.source_type,
            "account_origin": origin_for_account(a),
            "account_domain": domain_for_account(a) or "",
            "principal_source": a.principal_source or "",
            "principal_type": a.principal_type.value,
            "auth_source": a.auth_source.value,
            "enabled_status": a.enabled_status.value,
            "interactive_status": a.interactive_status.value,
            "last_login": a.last_login.isoformat() if a.last_login else "",
            "privilege_classification": a.privilege_classification.value,
            "privilege_confidence": a.privilege_confidence,
            "risk_score": a.risk_score,
            "is_shared": a.is_shared,
            "password_never_expires": a.password_never_expires,
            "owner": a.owner or "",
            "discovered_at": a.discovered_at.isoformat() if a.discovered_at else "",
        }
        for a in accounts
    ]


@router.get("/accounts/csv")
def export_accounts_csv(
    only_privileged: bool = False,
    limit: int = Query(default=_EXPORT_MAX_ROWS, ge=1, le=_EXPORT_MAX_ROWS),
    offset: int = Query(default=0, ge=0),
    db: Session = Depends(get_db),
    p: Principal = Depends(require_roles("admin", "security_analyst", "auditor")),
):
    import csv

    q = db.query(Account)
    if only_privileged:
        q = q.filter(Account.privilege_classification.in_([c.value for c in PRIVILEGED_CLASSES]))
    accounts = q.order_by(Account.risk_score.desc()).offset(offset).limit(limit).all()
    rows = _accounts_to_rows(accounts)

    buf = io.StringIO()
    if rows:
        writer = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)

    log_action(db, "export.accounts.csv", actor_id=p.id, actor_email=p.email, context={"only_privileged": only_privileged, "count": len(rows)})
    db.commit()

    buf.seek(0)
    filename = f"accounts_{'privileged_' if only_privileged else ''}export_{datetime.now(UTC).strftime('%Y%m%d_%H%M%S')}.csv"
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/accounts/excel")
def export_accounts_excel(
    only_privileged: bool = False,
    limit: int = Query(default=_EXPORT_MAX_ROWS, ge=1, le=_EXPORT_MAX_ROWS),
    offset: int = Query(default=0, ge=0),
    db: Session = Depends(get_db),
    p: Principal = Depends(require_roles("admin", "security_analyst", "auditor")),
):
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    q = db.query(Account)
    if only_privileged:
        q = q.filter(Account.privilege_classification.in_([c.value for c in PRIVILEGED_CLASSES]))
    accounts = q.order_by(Account.risk_score.desc()).offset(offset).limit(limit).all()
    rows = _accounts_to_rows(accounts)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Accounts"
    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)
        header_font = Font(bold=True)
        for cell in ws[1]:
            cell.font = header_font
        for row in rows:
            ws.append(list(row.values()))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    log_action(db, "export.accounts.excel", actor_id=p.id, actor_email=p.email, context={"only_privileged": only_privileged, "count": len(rows)})
    db.commit()

    filename = f"accounts_{'privileged_' if only_privileged else ''}export_{datetime.now(UTC).strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        io.BytesIO(buf.getvalue()),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


_HOUSEKEEPING_STATUSES = [
    ActivityStatus.inactive_30d,
    ActivityStatus.inactive_90d,
    ActivityStatus.never_logged_in,
    ActivityStatus.no_evidence,
]


@router.get("/housekeeping/csv")
def export_housekeeping_csv(
    limit: int = Query(default=_EXPORT_MAX_ROWS, ge=1, le=_EXPORT_MAX_ROWS),
    offset: int = Query(default=0, ge=0),
    db: Session = Depends(get_db),
    p: Principal = Depends(require_roles("admin", "security_analyst", "auditor")),
):
    """Monthly housekeeping report: enabled accounts with stale or missing login evidence.

    Risk ratings follow the governance matrix (30d=medium, 90d=high,
    privileged 30d=high, privileged 90d=critical, no evidence=review required).
    """
    import csv

    accounts = (
        db.query(Account)
        .filter(Account.enabled_status == EnabledStatus.enabled)
        .filter(Account.activity_status.in_(_HOUSEKEEPING_STATUSES))
        .order_by(Account.last_login.asc().nullsfirst())
        .offset(offset)
        .limit(limit)
        .all()
    )

    rows = []
    for a in accounts:
        rating, action = housekeeping_risk(a.activity_status, a.privilege_classification)
        rows.append({
            "hostname": a.asset.hostname if a.asset else "",
            "ip_address": (a.asset.ip_address or "") if a.asset else "",
            "platform": a.platform.value,
            "account_name": a.account_name,
            "principal_type": a.principal_type.value,
            "privilege_classification": a.privilege_classification.value,
            "activity_status": a.activity_status.value,
            "last_login": a.last_login.isoformat() if a.last_login else "",
            "last_login_source": a.last_login_source or "",
            "never_logged_in": a.never_logged_in if a.never_logged_in is not None else "",
            "password_never_expires": a.password_never_expires,
            "password_last_changed": a.password_last_changed.isoformat() if a.password_last_changed else "",
            "owner": a.owner or "",
            "risk_rating": rating,
            "recommended_action": action,
            "collection_mode": a.collection_mode or "",
            "discovered_at": a.discovered_at.isoformat() if a.discovered_at else "",
        })

    buf = io.StringIO()
    fieldnames = list(rows[0].keys()) if rows else [
        "hostname", "ip_address", "platform", "account_name", "principal_type",
        "privilege_classification", "activity_status", "last_login", "last_login_source",
        "never_logged_in", "password_never_expires", "password_last_changed", "owner",
        "risk_rating", "recommended_action", "collection_mode", "discovered_at",
    ]
    writer = csv.DictWriter(buf, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerows(rows)

    log_action(db, "export.housekeeping.csv", actor_id=p.id, actor_email=p.email,
               context={"count": len(rows)})
    db.commit()

    buf.seek(0)
    filename = f"housekeeping_{datetime.now(UTC).strftime('%Y%m%d_%H%M%S')}.csv"
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
